"""
AutoFlow Studio - Excel Pivot & Automated Reporting
Compatible with Streamlit Cloud, Vercel Python Serverless, and Local CLI
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
import sys

def generate_sales_report(input_file: str, output_file: str = None, report_title: str = "Sales Performance Report"):
    """
    Reads sales dataset, creates a pivot table, applies Excel formulas, styles headers,
    and embeds a native BarChart.
    """
    if not os.path.exists(input_file):
        raise FileNotFoundError(f"Input file '{input_file}' not found.")
        
    if output_file is None:
        base, ext = os.path.splitext(input_file)
        output_file = f"{base}_report{ext}"

    # Step 1: Create Pivot Table using Pandas
    df = pd.read_excel(input_file)
    pivot = df.pivot_table(
        index='Gender',
        columns='Product line',
        values='Total',
        aggfunc='sum'
    ).round(2).fillna(0)
    
    # Save base pivot to target file
    pivot.to_excel(output_file, startrow=4, sheet_name='Report')
    
    # Step 2: Open with openpyxl for formatting and charts
    wb = load_workbook(output_file)
    ws = wb['Report']
    
    min_col = ws.min_column
    max_col = ws.max_column
    min_row = 5
    max_row = ws.max_row
    
    # Step 3: Add Title & Metadata Header
    ws['A1'] = report_title
    ws['A1'].font = Font(name='Arial', size=16, bold=True, color='1E3A8A')
    ws['A2'] = "Automated by AutoFlow Studio | Streamlit & Cloud Ready"
    ws['A2'].font = Font(name='Arial', size=9, italic=True, color='64748B')
    
    # Step 4: Inject =SUM() Formulas at bottom of columns
    total_row = max_row + 1
    ws[f'A{total_row}'] = "Total"
    ws[f'A{total_row}'].font = Font(name='Arial', bold=True, color='1E3A8A')
    
    for c in range(min_col + 1, max_col + 1):
        col_letter = get_column_letter(c)
        sum_cell = ws[f'{col_letter}{total_row}']
        sum_cell.value = f"=SUM({col_letter}{min_row}:{col_letter}{max_row})"
        sum_cell.font = Font(bold=True)
        sum_cell.number_format = '$#,##0.00'
        
    # Step 5: Format Data Cells with Currency Format
    for r in range(min_row, max_row + 1):
        for c in range(min_col + 1, max_col + 1):
            col_letter = get_column_letter(c)
            ws[f'{col_letter}{r}'].number_format = '$#,##0.00'
            
    # Step 6: Create & Embed BarChart
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Sales Total by Product Line"
    chart.y_axis.title = "Total USD"
    chart.x_axis.title = "Demographic"
    
    data = Reference(ws, min_col=min_col + 1, max_col=max_col, min_row=4, max_row=max_row)
    categories = Reference(ws, min_col=min_col, max_col=min_col, min_row=min_row, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.width = 16
    chart.height = 10
    
    ws.add_chart(chart, f"B{total_row + 3}")
    
    wb.save(output_file)
    print(f"Report generated successfully: {output_file}")
    return output_file

if __name__ == "__main__":
    sample = "sales_february.xlsx"
    if os.path.exists(sample):
        generate_sales_report(sample)
    else:
        print("Please provide a valid Excel sales file path.")
